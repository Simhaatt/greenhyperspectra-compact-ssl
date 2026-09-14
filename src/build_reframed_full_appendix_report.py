from __future__ import annotations
import os

import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt

ROOT = Path(os.environ.get("HSI_ROOT", Path(__file__).resolve().parents[1]))
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from build_reframed_detailed_report import (  # noqa: E402
    FIG_DIR,
    PAPER_FIG_DIR,
    RESULTS,
    TABLE_DIR,
    add_df_table,
    add_image,
    landscape_section,
    portrait_section,
)


BASE = RESULTS / "detailed_report_reframed_complete.docx"
OUT = RESULTS / "detailed_report_reframed_full_appendix.docx"


def table_title(path: Path) -> str:
    name = path.stem.replace("_", " ")
    return "Appendix Table. " + name[:1].upper() + name[1:]


def figure_title(path: Path) -> str:
    name = path.stem.replace("_", " ")
    return "Appendix Figure. " + name[:1].upper() + name[1:]


def add_intro(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    if p.runs:
        p.runs[0].font.size = Pt(9)


def csv_paths() -> list[Path]:
    paths: list[Path] = []
    folders = [
        RESULTS / "paper_tables",
        RESULTS / "summary",
        RESULTS / "band_selection",
        RESULTS / "rl",
        RESULTS / "beam",
        RESULTS / "beam_stability",
        RESULTS / "beam_supervised_modern_heads",
        RESULTS / "hybrid_ssl_label_efficiency",
        RESULTS / "ood_source_validation",
        RESULTS / "source_recalibration",
    ]
    seen: set[Path] = set()
    for folder in folders:
        if not folder.exists():
            continue
        for path in sorted(folder.glob("*.csv")):
            resolved = path.resolve()
            if resolved not in seen:
                paths.append(path)
                seen.add(resolved)
    return paths


def figure_paths() -> list[Path]:
    paths: list[Path] = []
    seen: set[Path] = set()
    for folder in [PAPER_FIG_DIR, FIG_DIR]:
        if not folder.exists():
            continue
        for path in sorted(folder.glob("*.png")):
            resolved = path.resolve()
            if resolved not in seen:
                paths.append(path)
                seen.add(resolved)
    return paths


def compact_csv(df: pd.DataFrame) -> pd.DataFrame:
    # Keep appendices complete enough to audit while avoiding pathological Word tables.
    if len(df) > 120:
        return df.head(120).copy()
    return df


def main() -> None:
    if not BASE.exists():
        raise FileNotFoundError(BASE)
    doc = Document(BASE)

    landscape_section(doc)
    doc.add_heading("Appendix B. Complete Result Tables Archive", level=1)
    add_intro(
        doc,
        "This appendix restores the exhaustive archive style of the earlier long report. "
        "It includes generated CSV tables from paper_tables, summary outputs, band-selection runs, RL selection, BeamSearch outputs, band-stability outputs, label-efficiency, OOD validation, and few-sample source recalibration. "
        "Very long CSV files are shown with their first 120 rows to keep the Word document stable; the complete CSV files remain stored in the results folders.",
    )
    for path in csv_paths():
        try:
            df = pd.read_csv(path)
        except Exception as exc:
            doc.add_paragraph(f"Skipped {path.name}: {exc}")
            continue
        shown = compact_csv(df)
        note = f"Source file: {path.relative_to(ROOT)}; displayed rows: {len(shown)} of {len(df)}; columns: {len(df.columns)}."
        add_df_table(doc, shown, table_title(path), note=note, max_rows=None, digits=4)

    portrait_section(doc)
    doc.add_heading("Appendix C. Complete Figure Archive", level=1)
    add_intro(
        doc,
        "This appendix embeds the generated paper figures and supporting experiment figures, including band selection, SSL, hybrid heads, label-efficiency, and OOD/source-validation plots.",
    )
    for path in figure_paths():
        add_image(doc, path, figure_title(path), width=6.25)

    doc.save(OUT)
    print(OUT)
    print("tables", len(doc.tables))
    print("figures", len(doc.inline_shapes))
    print("csv_files", len(csv_paths()))
    print("figure_files", len(figure_paths()))


if __name__ == "__main__":
    main()
