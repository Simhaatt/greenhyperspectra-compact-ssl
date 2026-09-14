from __future__ import annotations
import os

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(os.environ.get("HSI_ROOT", Path(__file__).resolve().parents[1]))
REPORTS = [
    (
        ROOT / "results" / "experiment_summary.md",
        ROOT / "results" / "experiment_summary.docx",
        "Experiment Summary",
        "GreenHyperSpectra compact wavelength regression",
    ),
    (
        ROOT / "results" / "detailed_report.md",
        ROOT / "results" / "detailed_report.docx",
        "Detailed Project Report",
        "Band-efficient hyperspectral plant trait regression",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, bottom=80, start=120, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for tag, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(subtitle)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string("555555")


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("222222")


def is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped.startswith("|") or not stripped.endswith("|"):
        return False
    cells = [cell.strip() for cell in stripped.strip("|").split("|")]
    return all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int] | None:
    if start + 1 >= len(lines) or not is_table_separator(lines[start + 1]):
        return None
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines):
        line = lines[idx].strip()
        if not (line.startswith("|") and line.endswith("|")):
            break
        if not is_table_separator(line):
            rows.append([clean_inline(cell.strip()) for cell in line.strip("|").split("|")])
        idx += 1
    return rows, idx


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)
    set_cell_margins(table)

    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            text = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            if r_idx == 0:
                run.font.bold = True
                set_cell_shading(cell, "F2F4F7")
    doc.add_paragraph()


def add_image(doc: Document, alt_text: str, image_path: str) -> None:
    path = Path(image_path)
    if not path.is_absolute():
        path = ROOT / image_path
    if not path.exists():
        add_rich_paragraph(doc, f"[Missing figure: {image_path}]")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.6))
    if alt_text:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)
        caption_run = caption.add_run(alt_text)
        caption_run.font.name = "Calibri"
        caption_run.font.size = Pt(9)
        caption_run.font.italic = True
        caption_run.font.color.rgb = RGBColor.from_string("555555")


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text


def add_rich_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(token)


def add_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    idx = 0
    in_code = False
    code_lines: list[str] = []

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if not stripped:
            idx += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            add_image(doc, clean_inline(image_match.group(1)), image_match.group(2))
            idx += 1
            continue

        if stripped.startswith("|"):
            parsed = parse_table(lines, idx)
            if parsed:
                rows, idx = parsed
                add_table(doc, rows)
                continue

        if stripped.startswith("# "):
            add_rich_paragraph(doc, clean_inline(stripped[2:]), "Heading 1")
        elif stripped.startswith("## "):
            add_rich_paragraph(doc, clean_inline(stripped[3:]), "Heading 2")
        elif stripped.startswith("### "):
            add_rich_paragraph(doc, clean_inline(stripped[4:]), "Heading 3")
        elif stripped.startswith("- "):
            add_rich_paragraph(doc, clean_inline(stripped[2:]), "List Bullet")
        elif re.match(r"^\d+\.\s+", stripped):
            add_rich_paragraph(doc, clean_inline(re.sub(r"^\d+\.\s+", "", stripped)), "List Number")
        else:
            add_rich_paragraph(doc, stripped)
        idx += 1

    if code_lines:
        add_code_block(doc, code_lines)


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Generated from project reports")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("777777")


def build_docx(source: Path, output: Path, title: str, subtitle: str) -> None:
    doc = Document()
    style_document(doc)
    add_title(doc, title, subtitle)
    markdown = source.read_text(encoding="utf-8")
    # The title is already rendered as a Word title block.
    markdown = re.sub(r"^# .+?$", "", markdown, count=1, flags=re.MULTILINE).lstrip()
    add_markdown(doc, markdown)
    add_footer(doc)
    doc.save(output)


def main() -> None:
    for source, output, title, subtitle in REPORTS:
        build_docx(source, output, title, subtitle)
        print(output)


if __name__ == "__main__":
    main()
