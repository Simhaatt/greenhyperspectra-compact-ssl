from __future__ import annotations
import os

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(os.environ.get("HSI_ROOT", Path(__file__).resolve().parents[1]))
RESULTS = ROOT / "results"
TABLE_DIR = RESULTS / "paper_tables"
FIG_DIR = RESULTS / "figures"
PAPER_FIG_DIR = RESULTS / "paper_figures"
OUT = RESULTS / "detailed_report_reframed_complete.docx"
OUT_V2 = RESULTS / "detailed_report_reframed_complete_with_band_selection.docx"
OUT_V3 = RESULTS / "detailed_report_reframed_complete_with_fusion_ablation.docx"
MD_OUT = RESULTS / "detailed_report_reframed_complete.md"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 38, 56)
MUTED = RGBColor(89, 89, 89)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
LIGHT_GREEN = "EAF7EA"
LIGHT_GOLD = "FFF6D9"


def read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(path)
    return pd.read_csv(path)


def fmt(value: object, digits: int = 3) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        value = float(value)
        if abs(value) >= 1000 or (0 < abs(value) < 0.001):
            return f"{value:.3e}"
        return f"{value:.{digits}f}".rstrip("0").rstrip(".")
    return str(value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, bottom=80, start=120, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for tag, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, width_dxa: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Detailed Project Report")
    r.font.name = "Calibri"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(
        "Band-efficient hyperspectral plant trait prediction using labelled calibration, "
        "self-supervised spectral context, and source-shift validation"
    )
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    run = meta.add_run("Dataset: GreenHyperSpectra | Targets: 8 plant traits | Bands: 1721 | Unlabelled spectra: 139,295")
    run.bold = True
    run.font.color.rgb = INK


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)
    set_cell_margins(table, top=120, bottom=120, start=160, end=160)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body)
    doc.add_paragraph()


def add_df_table(
    doc: Document,
    df: pd.DataFrame,
    title: str,
    note: str | None = None,
    max_rows: int | None = None,
    digits: int = 3,
    widths: list[int] | None = None,
) -> None:
    doc.add_heading(title, level=2)
    if note:
        p = doc.add_paragraph(note)
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = MUTED
    shown = df.copy()
    if max_rows is not None:
        shown = shown.head(max_rows)
    table = doc.add_table(rows=1, cols=len(shown.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_table_width(table)
    set_cell_margins(table)
    set_repeat_table_header(table.rows[0])

    if widths and len(widths) == len(shown.columns):
        for idx, width in enumerate(widths):
            set_cell_width(table.rows[0].cells[idx], width)

    for idx, col in enumerate(shown.columns):
        cell = table.rows[0].cells[idx]
        if widths and len(widths) == len(shown.columns):
            set_cell_width(cell, widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(col))
        run.bold = True
        run.font.size = Pt(8)

    for _, row in shown.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(shown.columns):
            cell = cells[idx]
            if widths and len(widths) == len(shown.columns):
                set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            text = fmt(row[col], digits=digits)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 20 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(8 if len(shown.columns) <= 7 else 7)
    doc.add_paragraph()


def add_image(doc: Document, path: Path, title: str, width: float = 6.35) -> None:
    if not path.exists():
        return
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(f"Figure source: {path.name}")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)
    cap.runs[0].font.color.rgb = MUTED


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def landscape_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)


def portrait_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def labelled_head_table() -> pd.DataFrame:
    df = read_csv(RESULTS / "beam_supervised_modern_heads" / "beam_supervised_modern_heads_labelled_modern_heads_all8_3seeds_summary.csv")
    mean = df.groupby("model", as_index=False).agg(mean_r2=("r2_mean", "mean"), mean_rmse=("rmse_mean", "mean"))
    return mean.sort_values("mean_r2", ascending=False).round(4)


def hybrid_leaderboard() -> pd.DataFrame:
    frames = []
    files = [
        RESULTS / "hybrid_ssl_beam_gated_heads" / "hybrid_ssl_beam_gated_heads_cross_gated_interaction_variants_all8_3seeds_summary.csv",
        RESULTS / "hybrid_ssl_beam_gated_heads" / "hybrid_ssl_beam_gated_heads_cross_gated_advanced_variants_all8_3seeds_summary.csv",
        RESULTS / "hybrid_ssl_beam_gated_heads" / "hybrid_ssl_beam_gated_heads_cross_gated_hyper_multiscale_all8_3seeds_summary.csv",
    ]
    for file in files:
        if file.exists():
            frames.append(pd.read_csv(file))
    df = pd.concat(frames, ignore_index=True)
    mean = df.groupby("model", as_index=False).agg(mean_r2=("r2_mean", "mean"), mean_rmse=("rmse_mean", "mean"))
    return mean.sort_values("mean_r2", ascending=False).drop_duplicates("model").head(10).round(4)


def best_overall_table() -> pd.DataFrame:
    path = TABLE_DIR / "table25_best_overall_model_per_trait.csv"
    df = read_csv(path)
    keep = [c for c in ["target", "best_family", "best_model", "best_r2"] if c in df.columns]
    if keep:
        df = df[keep]
    return df.round(4)


def label_efficiency_table() -> pd.DataFrame:
    path = TABLE_DIR / "table32_hybrid_ssl_label_efficiency_mean_r2_by_fraction.csv"
    df = read_csv(path)
    return df.round(4)


def fusion_ablation_leaderboard() -> pd.DataFrame:
    path = TABLE_DIR / "table37_fusion_ablation_leaderboard.csv"
    if not path.exists():
        return pd.DataFrame()
    df = read_csv(path)
    cols = [c for c in ["fusion_method", "model", "mean_r2", "best_traits", "comment"] if c in df.columns]
    return df[cols].round(4)


def fusion_ablation_best_per_trait() -> pd.DataFrame:
    path = TABLE_DIR / "table39_fusion_ablation_best_per_trait.csv"
    if not path.exists():
        return pd.DataFrame()
    df = read_csv(path)
    cols = [c for c in ["target", "fusion_method", "model", "r2_mean", "rmse_mean", "r2_std"] if c in df.columns]
    return df[cols].round(4)


def ood_model_table() -> pd.DataFrame:
    path = RESULTS / "summary" / "table35_ood_source_validation_model_leaderboard.csv"
    df = read_csv(path)
    cols = [c for c in ["family", "model", "mean_r2", "median_r2", "std_r2", "mean_rpd", "n"] if c in df.columns]
    return df[cols].round(4)


def ood_trait_table() -> pd.DataFrame:
    path = RESULTS / "summary" / "table36_ood_source_validation_trait_model_summary.csv"
    df = read_csv(path)
    idx = df.groupby("target")["mean_r2"].idxmax()
    out = df.loc[idx, ["target", "model", "mean_r2", "std", "n"]].sort_values("target")
    return out.round(4)


def significance_table() -> pd.DataFrame:
    path = TABLE_DIR / "table27_paired_trait_significance_tests.csv"
    df = read_csv(path)
    cols = [c for c in ["comparison", "mean_diff_r2", "bootstrap_ci_low", "bootstrap_ci_high", "paired_sign_p"] if c in df.columns]
    return df[cols].round(4)


def statistical_model_comparison_table() -> pd.DataFrame:
    path = TABLE_DIR / "table40_statistical_model_comparison.csv"
    if not path.exists():
        return pd.DataFrame()
    df = read_csv(path)
    cols = [
        c
        for c in [
            "comparison",
            "mean_delta_r2",
            "ci95_low",
            "ci95_high",
            "wilcoxon_p_value",
            "traits_improved",
            "traits_compared",
            "interpretation",
        ]
        if c in df.columns
    ]
    return df[cols].round(4)


def band_stability_table() -> pd.DataFrame:
    path = TABLE_DIR / "table41_band_stability_analysis.csv"
    if not path.exists():
        return pd.DataFrame()
    df = read_csv(path)
    cols = [
        c
        for c in [
            "trait",
            "runs",
            "mean_exact_band_overlap",
            "mean_exact_jaccard",
            "mean_region_overlap",
            "mean_region_jaccard",
            "most_frequent_regions",
        ]
        if c in df.columns
    ]
    return df[cols].round(4)


def source_recalibration_table() -> pd.DataFrame:
    extended_path = TABLE_DIR / "table43_few_sample_source_recalibration_extended.csv"
    path = extended_path if extended_path.exists() else TABLE_DIR / "table42_few_sample_source_recalibration.csv"
    if not path.exists():
        return pd.DataFrame()
    df = read_csv(path)
    cols = [
        c
        for c in [
            "model",
            "target_calibration_percent",
            "mean_r2",
            "median_r2",
            "rmse",
            "rpd",
            "traits_improved",
            "n_evaluations",
        ]
        if c in df.columns
    ]
    return df[cols].round(4)


def band_selection_top_table() -> pd.DataFrame:
    df = read_csv(TABLE_DIR / "table3_band_selection_comparison.csv")
    cols = ["target", "method", "budget", "rmse", "r2", "rpd"]
    df = df[cols].copy()
    df = df.sort_values(["target", "r2"], ascending=[True, False])
    return df.groupby("target", as_index=False).head(8).round(4)


def compact_primary_table() -> pd.DataFrame:
    df = read_csv(TABLE_DIR / "table4_compact_model_comparison.csv")
    return df.round(4)


def rl_summary_table() -> pd.DataFrame:
    frames = []
    for target in ["cab", "cw", "cm"]:
        path = RESULTS / "rl" / f"{target}_rl_band_selection_results.csv"
        if path.exists():
            df = pd.read_csv(path)
            if "target" not in df.columns:
                df.insert(0, "target", target)
            frames.append(df)
    if not frames:
        return pd.DataFrame()
    df = pd.concat(frames, ignore_index=True)
    keep = [c for c in ["target", "method", "budget", "rmse", "mae", "r2", "rpd", "selected_wavelengths"] if c in df.columns]
    return df[keep].round(4)


def interpretation_table() -> pd.DataFrame:
    path = TABLE_DIR / "table31_traitwise_wavelength_interpretation_summary.csv"
    df = read_csv(path)
    cols = [c for c in ["target", "visible_red_edge_count", "nir_count", "swir_count", "main_interpretation"] if c in df.columns]
    return df[cols]


def write_markdown_snapshot() -> None:
    text = """# Detailed Project Report - Reframed

This report reframes the work as labelled calibration plus unlabelled spectral representation learning. The labelled split is the source of trait supervision and evaluation. The unlabelled spectra are not treated as a competing dataset; they provide spectral context for self-supervised learning and can be assigned model-predicted traits after calibration.

The final experimental story is: compact supervised wavelength selection is the strongest in-distribution baseline; SSL-based full-spectrum context is a complementary extension; label-efficiency shows hybrid models are close but not consistently superior; OOD validation reveals strong source shift; few-sample source recalibration substantially recovers OOD performance; and band-stability analysis shows that selected spectral regions remain physiologically consistent even when exact wavelengths vary.
"""
    MD_OUT.write_text(text, encoding="utf-8")


def build() -> None:
    write_markdown_snapshot()
    doc = Document()
    style_document(doc)
    add_title(doc)

    add_callout(
        doc,
        "Correct framing for the project",
        "This project does not compare labelled data against unlabelled data. Labelled spectra provide the calibration signal: they teach the model how spectral patterns map to plant traits and they provide the held-out test evidence. Unlabelled spectra provide additional spectral context through self-supervised learning, so that the final calibrated model can be applied to large unlabelled collections as predicted trait estimates. Predicted traits on unlabelled spectra are model outputs, not new ground truth labels.",
        LIGHT_GREEN,
    )
    doc.add_paragraph(
        "In plain terms, the report is reframed around using labelled data to learn and validate the trait mapping, then using that calibrated mapping to estimate traits for unlabelled spectra. "
        "The experiments are therefore not a comparison of labelled data versus unlabelled data. They test whether unlabelled spectral context, learned through SSL and fused with labelled calibration, helps the final trait-prediction system."
    )

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "The project developed a complete pipeline for GreenHyperSpectra trait prediction from hyperspectral reflectance. "
        "The pipeline begins with labelled trait supervision, selects compact trait-specific wavelength subsets, trains modern regression heads, "
        "learns self-supervised full-spectrum context from 139,295 unlabelled spectra, and evaluates whether this context improves prediction under ordinary splits, low-label regimes, and source shift."
    )
    add_bullets(
        doc,
        [
            "Main calibrated model: BeamSearch K=30 wavelengths plus a residual MLP head trained on labelled rows.",
            "SSL-aware extension: 30 selected wavelengths plus a 128-dimensional embedding from a contiguous-mask self-supervised encoder.",
            "Fusion method: cross-gated Beam-SSL fusion, where selected wavelengths and full-spectrum SSL context modulate each other before regression.",
            "Best labelled-only family: BeamK30 ResidualMLP with mean R2 about 0.624 across eight traits.",
            "Best SSL-aware hybrid family: Cross-Gated NAM head with mean R2 about 0.598 to 0.601 depending on the run summary.",
            "Label-efficiency result: hybrid models are close at 50 percent and 100 percent labels but do not beat the strong labelled residual model on average.",
            "OOD result: zero-shot leave-source-out transfer is difficult, with negative mean R2 before source calibration.",
            "Few-sample source recalibration: adding 10 to 20 percent labelled target-source calibration samples turns mean OOD R2 positive for both BeamK30 ResidualMLP and Hybrid Cross-Gated NAM.",
            "Statistical comparison: Hybrid NAM, Additive, and FiLM fusion are close to the BeamK30 residual baseline, while SSL-only is clearly weaker.",
            "Band stability: exact selected wavelengths vary across resampling, but selected spectral regions remain relatively stable and physiologically consistent.",
        ],
    )

    doc.add_heading("2. Dataset and Supervision", level=1)
    doc.add_paragraph(
        "The labelled train/test files contain 5,635 labelled spectra after combining the official train and test splits. "
        "The modelling spectrum used in the experiments contains 1,721 bands from 400 nm to 2450 nm. "
        "The unlabelled HuggingFace dataset contains 139,295 spectra. These unlabelled rows do not contain measured Cab, Cw, Cm, LAI, Cp, Cbc, Car, or Anth values, so they cannot directly evaluate trait prediction."
    )
    add_df_table(
        doc,
        read_csv(TABLE_DIR / "table1_dataset_summary.csv"),
        "Table 1. Dataset summary",
        "The labelled split is used for supervised calibration and evaluation; the unlabelled split is used for representation learning and potential trait-map generation.",
    )
    doc.add_paragraph(
        "This distinction matters scientifically. Performance metrics such as R2 and RMSE can only be computed where measured labels exist. "
        "For unlabelled spectra, the trained model can generate predicted trait estimates, but those estimates remain predictions unless later validated by measurements."
    )

    doc.add_heading("3. End-to-End Process Followed", level=1)
    add_numbered(
        doc,
        [
            "Install and verify the GreenHyperSpectra labelled and unlabelled files.",
            "Train full-spectrum baselines on all 1,721 bands to establish reference performance.",
            "Run BeamSearch wavelength selection separately for each target to obtain compact K=30 wavelength sets.",
            "Train compact labelled models on the selected bands, including MLP, residual MLP, KAN, DCN, NAM, and polynomial heads.",
            "Pretrain a 1D convolutional masked autoencoder on unlabelled spectra, first with random masking and then with contiguous spectral-region masking.",
            "Extract a 128-dimensional SSL embedding from the pretrained encoder for labelled spectra.",
            "Fuse the 30 selected wavelengths with the 128-dimensional SSL embedding using cross-gated fusion and evaluate multiple heads.",
            "Run label-efficiency experiments from 1 percent to 100 percent labels to test whether SSL context helps when labelled measurements are scarce.",
            "Run leave-source-out validation using the dataset/source column to test domain shift.",
            "Generate paper tables, figures, Word reports, Kaggle notebooks, and downloadable result bundles.",
        ],
    )

    doc.add_heading("4. Methodology and Architecture", level=1)
    doc.add_heading("4.1 BeamSearch wavelength selection", level=2)
    doc.add_paragraph(
        "BeamSearch is used to select compact trait-specific wavelength subsets. At each step, candidate wavelengths are added to partial subsets and scored using validation performance. "
        "Keeping several partial subsets avoids the brittleness of greedy one-step selection while remaining practical for high-dimensional spectra."
    )
    doc.add_heading("4.2 Main labelled calibration model", level=2)
    doc.add_paragraph(
        "The strongest calibrated baseline uses 30 BeamSearch-selected wavelengths per trait and a residual MLP regression head. "
        "This is the primary in-distribution model because it uses measured labels directly, is compact enough for future multispectral deployment, and generalized better than several newer heads on the held-out labelled test split."
    )
    doc.add_heading("4.3 Self-supervised spectral context", level=2)
    doc.add_paragraph(
        "Self-supervised learning was used to learn how vegetation spectra are structured without requiring trait labels. "
        "The encoder is trained by masking wavelength bands and reconstructing the missing reflectance values. Contiguous masking was added because it is more realistic for spectra: the model must infer missing wavelength regions from surrounding spectral context rather than filling isolated random points."
    )
    doc.add_heading("4.4 Hybrid Beam-SSL model", level=2)
    doc.add_paragraph(
        "The SSL-aware model combines two information paths. Path A is compact and interpretable: 30 selected wavelengths from BeamSearch. Path B is contextual: a 128-dimensional embedding generated by the pretrained full-spectrum SSL encoder. "
        "The cross-gated fusion module learns how Beam features should condition SSL context and how SSL context should condition the selected wavelengths. This makes the hybrid model a calibrated labelled model enhanced by unlabelled spectral pretraining, not a replacement of labelled supervision."
    )
    add_bullets(
        doc,
        [
            "Beam branch: 30 selected wavelength features projected into a 64-dimensional representation.",
            "SSL branch: 128-dimensional full-spectrum encoder embedding projected into a 64-dimensional representation.",
            "Cross gating: each branch learns gates that modulate the other branch.",
            "Fusion: gated Beam and SSL representations are concatenated into a 128-dimensional fused representation.",
            "Regression heads tested: MLP, NAM, DCN, bilinear, highway, KAN-enhanced, residual, MoE, hypernetwork, multiscale, and physics-guided multitask variants.",
        ],
    )
    doc.add_heading("4.5 Using labelled calibration to predict unlabelled spectra", level=2)
    doc.add_paragraph(
        "After a trait model is calibrated on labelled spectra and validated on held-out labelled spectra, it can be applied to unlabelled spectra to produce estimated trait values. "
        "This is the practical deployment path: measured labels teach the mapping, unlabelled spectra expand coverage, and the model outputs predicted Cab, Cw, Cm, LAI, Cp, Cbc, Car, and Anth values where direct lab measurements are unavailable."
    )

    doc.add_heading("5. Results", level=1)
    doc.add_heading("5.1 Band-selection and compact-model ablations", level=2)
    doc.add_paragraph(
        "The report includes the earlier ablation path because it explains why BeamSearch K=30 became the compact feature extractor. "
        "PLSR was used as a fast scoring model during selection, while later neural heads were trained after the wavelength subsets had been chosen. "
        "The important comparison is therefore not only the final neural model, but the sequence of band selectors that led to the final compact representation."
    )
    add_df_table(
        doc,
        band_selection_top_table(),
        "Table 2. Band-selection comparison across PLSR, RL, random, correlation, mutual information, and KAN variants",
        "This table promotes the earlier selector results into the main report. BeamSearch+PLSR is retained because it is stable and consistently competitive, while RL-style selection is treated as a useful exploratory ablation.",
        max_rows=30,
        digits=4,
    )
    add_df_table(
        doc,
        compact_primary_table(),
        "Table 3. Compact model comparison for primary traits",
        "After band selection, compact neural regressors were trained on selected wavelengths. This is where Beam K=30 + TorchMLP/ResidualMLP became the strong supervised baseline.",
        digits=4,
    )
    rl_df = rl_summary_table()
    if not rl_df.empty:
        add_df_table(
            doc,
            rl_df,
            "Table 4. RL-style wavelength-selection summary",
            "The epsilon-greedy/RL-style selector was evaluated mainly as an exploratory sequential selection method. It was informative, but BeamSearch was more reliable for the final pipeline.",
            max_rows=12,
            digits=4,
        )

    doc.add_heading("5.2 Full-spectrum and final calibrated models", level=2)
    add_df_table(
        doc,
        read_csv(TABLE_DIR / "table2_full_spectrum_baselines.csv").head(12),
        "Table 5. Full-spectrum baseline sample",
        "Full-spectrum models establish the reference for using all 1,721 bands. The compact models are judged against these baselines.",
        max_rows=12,
    )
    add_df_table(
        doc,
        labelled_head_table(),
        "Table 6. Labelled-only modern head leaderboard",
        "Mean performance across eight traits on the labelled held-out split. This is supervised calibration performance, not unlabelled-vs-labelled comparison.",
    )
    add_df_table(
        doc,
        hybrid_leaderboard(),
        "Table 7. SSL-aware hybrid model leaderboard",
        "Hybrid models use labelled calibration plus SSL context learned from unlabelled spectra.",
    )
    add_df_table(
        doc,
        best_overall_table(),
        "Table 8. Best model per trait",
        "The best per-trait model can come from labelled-only, hybrid SSL, or physics-guided families.",
    )
    fusion_df = fusion_ablation_leaderboard()
    if not fusion_df.empty:
        doc.add_heading("5.3 Fusion ablation", level=2)
        doc.add_paragraph(
            "The fusion ablation isolates the role of the fusion mechanism. Beam-only tests the compact labelled baseline; SSL-only tests whether the 128-dimensional SSL embedding alone carries enough trait signal; concat, additive, and FiLM test simple fusion baselines; Cross-gated NAM tests the proposed hybrid fusion/head combination."
        )
        add_df_table(
            doc,
            fusion_df,
            "Table 9. Fusion ablation leaderboard",
            "Beam-only remains strongest on average, but the hybrid fusion variants substantially outperform naive concatenation. Additive fusion and Cross-gated NAM are close, which means the paper should claim improved fusion over concat rather than a universal win over Beam-only.",
            digits=4,
        )
        best_fusion = fusion_ablation_best_per_trait()
        if not best_fusion.empty:
            add_df_table(
                doc,
                best_fusion,
                "Table 10. Best fusion method per trait",
                "Hybrid fusion wins selected traits: FiLM for LAI, additive for Cab, and Cross-gated NAM for Car. Beam-only remains best for Anth, Cbc, Cm, Cp, and Cw.",
                digits=4,
            )

    doc.add_heading("5.4 Label-efficiency, source shift, recalibration, and statistical checks", level=2)
    add_df_table(
        doc,
        label_efficiency_table(),
        "Table 11. Hybrid SSL label-efficiency mean R2 by fraction",
        "This evaluates how much labelled supervision is needed after SSL pretraining. The strong labelled residual model remains best on average, but the hybrid NAM model becomes close at higher label fractions.",
    )
    add_df_table(
        doc,
        ood_model_table(),
        "Table 12. Leave-source-out OOD model leaderboard",
        "OOD/source validation holds out entire dataset/source groups. Negative R2 values indicate strong source shift, not a failed run.",
    )
    add_df_table(
        doc,
        ood_trait_table(),
        "Table 13. Best leave-source-out model per trait",
        "Hybrid NAM is best for LAI, cp, and cw under source shift, while the labelled residual model remains best for cab, car, cbc, and cm.",
    )
    recalibration_df = source_recalibration_table()
    if not recalibration_df.empty:
        add_df_table(
            doc,
            recalibration_df,
            "Table 14. Few-sample source recalibration",
            "A simple prediction-level linear calibration layer, y_calibrated = a*y_pred + b, is fitted on a small labelled subset from each held-out source. The zero-shot 0 percent rows show the difficult source-shift baseline; 10 to 20 percent target-source calibration recovers positive mean OOD R2 for both tested models.",
            digits=4,
        )
    statistical_df = statistical_model_comparison_table()
    if not statistical_df.empty:
        add_df_table(
            doc,
            statistical_df,
            "Table 15. Statistical comparison against BeamK30 ResidualMLP",
            "Paired seed-by-trait comparisons support the competitive framing: hybrid fusion variants are close to the compact supervised model, whereas SSL-only is clearly weaker.",
            digits=4,
        )
    stability_df = band_stability_table()
    if not stability_df.empty:
        add_df_table(
            doc,
            stability_df,
            "Table 16. BeamSearch K=30 band stability",
            "Exact selected wavelengths vary across repeated BeamSearch runs, but region-level overlap remains stronger and supports the physiological consistency of selected spectral regions.",
            digits=4,
        )
    add_df_table(
        doc,
        significance_table(),
        "Table 17. Earlier bootstrap and paired-sign test summary",
        "These tests reduce overclaiming by showing whether model differences are stable across traits.",
    )
    add_df_table(
        doc,
        interpretation_table(),
        "Table 18. Trait-wise wavelength interpretation summary",
        "Selected bands align with expected visible/red-edge, NIR, and SWIR regions for pigments, structure, water, and biochemical traits.",
    )

    doc.add_heading("6. Interpretation of New Experiments", level=1)
    add_callout(
        doc,
        "Label-efficiency interpretation",
        "The label-efficiency experiment does not show a broad average win for the hybrid SSL model. Instead, it shows that the compact supervised residual model is already highly label-efficient. The hybrid model is scientifically useful because it stays close while incorporating full-spectrum unlabelled context, and it wins selected trait/fraction cases.",
        LIGHT_GOLD,
    )
    add_callout(
        doc,
        "OOD/source-shift interpretation",
        "The leave-source-out experiment is much harder than the normal held-out split. Zero-shot source transfer gives negative mean R2, revealing strong source shift. Few-sample source recalibration directly addresses this weakness: with 10 to 20 percent labelled calibration samples from the held-out source, both BeamK30 ResidualMLP and Hybrid Cross-Gated NAM recover positive mean OOD R2.",
        LIGHT_BLUE,
    )
    add_callout(
        doc,
        "Band-stability interpretation",
        "The repeated BeamSearch analysis shows the expected pattern for hyperspectral data: exact wavelengths are not identical across resampling, but the selected spectral regions remain much more stable. This supports an interpretability claim at the spectral-region level rather than overclaiming that individual wavelength indices are fixed.",
        LIGHT_GREEN,
    )
    doc.add_paragraph(
        "The correct conclusion is therefore nuanced: the final calibrated compact model is the strongest average in-distribution predictor, while the SSL-aware hybrid model is the proposed extension for using unlabelled spectra and for exploring robustness under source shift. "
        "The new recalibration result turns the OOD section from a weakness into a deployment story: zero-shot transfer is hard, but a small amount of target-source calibration can recover useful performance. This gives a better paper narrative than claiming that unlabelled data should outperform labelled data. Unlabelled data cannot provide trait supervision by itself; it improves the learned representation and broadens deployable coverage after calibration."
    )

    doc.add_heading("7. Safe Claims", level=1)
    add_bullets(
        doc,
        [
            "A 30-band BeamSearch representation can approach or exceed full-spectrum baselines for several traits.",
            "Residual MLP is the strongest average labelled-only compact head among the tested modern heads.",
            "Self-supervised pretraining on 139,295 unlabelled spectra learns reusable spectral context, represented as a 128-dimensional embedding.",
            "Cross-gated fusion provides an interpretable way to combine selected wavelengths with full-spectrum SSL context.",
            "Hybrid SSL models are competitive with, but not consistently better than, the strongest labelled-only baseline on the in-distribution split.",
            "Under zero-shot source shift, all methods struggle; with few-sample target-source recalibration, both tested models recover positive mean OOD R2.",
            "Exact BeamSearch-selected wavelengths vary across resampling, but selected spectral regions are stable enough to support region-level physiological interpretation.",
            "Predictions on unlabelled spectra should be described as estimated or pseudo trait values, not measured labels.",
        ],
    )

    doc.add_heading("8. Limitations", level=1)
    add_bullets(
        doc,
        [
            "The labelled dataset is modest compared with the dimensionality and diversity of hyperspectral spectra.",
            "OOD/source validation shows strong domain shift across dataset sources before target-source recalibration.",
            "Anthocyanin has only one labelled source in the source audit, so leave-source-out OOD validation is not meaningful for anth.",
            "The current SSL encoder improves context but does not guarantee average R2 improvement over the strongest labelled model.",
            "Unlabelled predictions need independent measurements before they can be treated as ground truth.",
            "Future compact deployment still needs a student model or sensor design that uses only selected wavelengths without requiring full-spectrum SSL embeddings at inference time.",
        ],
    )

    doc.add_heading("9. Recommended Final Paper Framing", level=1)
    doc.add_paragraph(
        "Recommended title: Label-Calibrated Hybrid Spectral Learning for Band-Efficient Plant Trait Prediction from GreenHyperSpectra."
    )
    doc.add_paragraph(
        "Recommended central claim: labelled measurements calibrate the trait mapping, while unlabelled spectra supply self-supervised spectral context that can be fused with compact selected wavelengths and used to generate trait estimates at scale. "
        "The strongest compact supervised model remains the best average in-distribution predictor, but the SSL-aware hybrid model provides a principled extension for unlabelled spectral collections. Source-shift experiments show that zero-shot transfer is difficult, while few-sample recalibration provides a practical deployment path."
    )
    doc.add_heading("10. What Next", level=1)
    add_numbered(
        doc,
        [
            "Update the final paper manuscript around the calibrated-plus-unlabelled-context framing used in this report.",
            "Generate a clean unlabelled trait-prediction export after selecting the final calibrated model family.",
            "Add a domain-adaptation section if more time is available, because OOD results show source shift is the main remaining weakness.",
            "Consider distilling the hybrid model into a 30-band student so compact deployment does not require full-spectrum SSL embeddings.",
            "Use the OOD and label-efficiency results honestly as evidence of both capability and limitations.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("11. Key Figures", level=1)
    figures = [
        (PAPER_FIG_DIR / "fig1_selected_wavelengths_annotated.png", "Figure 1. Selected wavelengths annotated by absorption region"),
        (FIG_DIR / "final_model_mean_r2_leaderboard.png", "Figure 2. Final model-family mean R2 leaderboard"),
        (FIG_DIR / "final_best_family_by_trait_r2.png", "Figure 3. Best model family by trait"),
        (FIG_DIR / "beam_supervised_modern_heads_labelled_modern_heads_all8_3seeds_r2.png", "Figure 4. Labelled modern heads R2"),
        (FIG_DIR / "hybrid_ssl_beam_gated_heads_cross_gated_advanced_variants_all8_3seeds_r2.png", "Figure 5. Cross-gated advanced hybrid variants R2"),
        (FIG_DIR / "hybrid_ssl_label_efficiency_mean_r2_all_traits.png", "Figure 6. Hybrid SSL label-efficiency mean R2"),
        (FIG_DIR / "hybrid_ssl_label_efficiency_delta_heatmap.png", "Figure 7. Hybrid SSL delta versus labelled baseline"),
        (FIG_DIR / "hybrid_ssl_fusion_ablation_fusion_ablation_all8_3seeds_leaderboard.png", "Figure 8. Fusion ablation leaderboard"),
        (FIG_DIR / "hybrid_ssl_fusion_ablation_fusion_ablation_all8_3seeds_by_trait.png", "Figure 9. Fusion ablation by trait"),
        (FIG_DIR / "ood_source_validation_top5_sources_3seeds_leaderboard.png", "Figure 10. OOD/source validation leaderboard"),
        (FIG_DIR / "ood_source_validation_top5_sources_3seeds_r2_by_trait.png", "Figure 11. OOD/source validation R2 by trait"),
    ]
    for path, title in figures:
        add_image(doc, path, title)

    landscape_section(doc)
    doc.add_heading("Appendix A. Additional Paper Tables", level=1)
    appendix_tables = [
        ("table3_band_selection_comparison.csv", "Appendix Table A1. Complete band-selection comparison"),
        ("table4_compact_model_comparison.csv", "Appendix Table A2. Compact model comparison"),
        ("table5_stability_results.csv", "Appendix Table A3. Beam/compact stability results"),
        ("table6_label_efficiency.csv", "Appendix Table A4. Original label-efficiency results"),
        ("table7_ssl_ablation_results.csv", "Appendix Table A5. SSL ablation results"),
        ("table8_wavelength_region_counts.csv", "Appendix Table A6. Wavelength-region counts"),
        ("table13_labelled_modern_heads_r2.csv", "Appendix Table A7. Labelled modern heads by trait"),
        ("table17_cross_gated_advanced_variants_r2.csv", "Appendix Table A8. Cross-gated advanced variants by trait"),
        ("table23_physics_multitask_r2.csv", "Appendix Table A9. Physics-guided multitask R2"),
        ("table24_labelled_hybrid_physics_final_comparison.csv", "Appendix Table A10. Labelled, hybrid, and physics comparison"),
        ("table32_hybrid_ssl_label_efficiency_mean_r2_by_fraction.csv", "Appendix Table A11. Hybrid label-efficiency mean R2"),
        ("table33_hybrid_ssl_label_efficiency_delta_vs_labelled.csv", "Appendix Table A12. Hybrid label-efficiency delta"),
        ("table37_fusion_ablation_leaderboard.csv", "Appendix Table A13. Fusion ablation leaderboard"),
        ("table38_fusion_ablation_trait_summary.csv", "Appendix Table A14. Fusion ablation trait summary"),
        ("table39_fusion_ablation_best_per_trait.csv", "Appendix Table A15. Fusion ablation best per trait"),
        ("table35_ood_source_validation_model_leaderboard.csv", "Appendix Table A16. OOD source validation model leaderboard"),
        ("table36_ood_source_validation_trait_model_summary.csv", "Appendix Table A17. OOD source validation by trait and model"),
        ("table40_statistical_model_comparison.csv", "Appendix Table A18. Statistical comparison versus BeamK30 ResidualMLP"),
        ("table41_band_stability_analysis.csv", "Appendix Table A19. BeamSearch K=30 band stability analysis"),
        ("table42_few_sample_source_recalibration.csv", "Appendix Table A20. Few-sample source recalibration"),
        ("table43_few_sample_source_recalibration_extended.csv", "Appendix Table A21. Extended few-sample source recalibration curve"),
    ]
    for name, title in appendix_tables:
        path = TABLE_DIR / name
        if path.exists():
            df = read_csv(path)
            add_df_table(doc, df, title, max_rows=40, digits=4)

    for candidate in [OUT, OUT_V2, OUT_V3]:
        try:
            doc.save(candidate)
            saved = candidate
            break
        except PermissionError:
            saved = None
    if saved is None:
        doc.save(RESULTS / "detailed_report_reframed_complete_latest.docx")
        saved = RESULTS / "detailed_report_reframed_complete_latest.docx"
    print(saved)
    print(MD_OUT)


if __name__ == "__main__":
    build()
