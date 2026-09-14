from __future__ import annotations
import os

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from markdown_reports_to_docx import build_docx


ROOT = Path(os.environ.get("HSI_ROOT", Path(__file__).resolve().parents[1]))
SOURCE = ROOT / "results" / "detailed_report.md"
OUTPUT = ROOT / "results" / "detailed_report.docx"
COMPLETE_OUTPUT = ROOT / "results" / "detailed_report_complete.docx"
TABLE_DIR = ROOT / "results" / "paper_tables"
PAPER_FIG_DIR = ROOT / "results" / "paper_figures"
FIG_DIR = ROOT / "results" / "figures"


TABLE_TITLES = {
    "table1_dataset_summary.csv": "Table 1. Dataset summary",
    "table2_full_spectrum_baselines.csv": "Table 2. Full-spectrum baselines",
    "table3_band_selection_comparison.csv": "Table 3. Band-selection comparison",
    "table4_compact_model_comparison.csv": "Table 4. Compact model comparison",
    "table5_stability_results.csv": "Table 5. Stability results",
    "table6_label_efficiency.csv": "Table 6. Label-efficiency results",
    "table6_label_efficiency_compact.csv": "Table 6b. Compact label-efficiency summary",
    "table7_ssl_ablation_results.csv": "Table 7. SSL ablation comparison",
    "table8_wavelength_region_counts.csv": "Table 8. Wavelength-region counts",
    "table13_labelled_modern_heads_r2.csv": "Table 13. Labelled-only modern heads R2",
    "table14_labelled_best_head_per_trait.csv": "Table 14. Best labelled-only head per trait",
    "table15_cross_gated_interaction_variants_r2.csv": "Table 15. Cross-gated interaction variants R2",
    "table16_best_labelled_vs_cross_gated_interaction.csv": "Table 16. Best labelled-only vs cross-gated interaction models",
    "table17_cross_gated_advanced_variants_r2.csv": "Table 17. Cross-gated advanced variants R2",
    "table18_best_labelled_vs_best_hybrid_so_far.csv": "Table 18. Best labelled-only vs best hybrid models",
    "table19_model_family_mean_r2_leaderboard.csv": "Table 19. Model-family mean R2 leaderboard",
    "table20_cross_gated_hyper_multiscale_r2.csv": "Table 20. Cross-gated hypernetwork and multiscale R2",
    "table21_best_labelled_vs_best_hybrid_after_hyper.csv": "Table 21. Best labelled-only vs best hybrid after hypernetwork",
    "table22_current_model_mean_r2_leaderboard.csv": "Table 22. Current model mean R2 leaderboard",
    "table23_physics_multitask_r2.csv": "Table 23. Physics-guided multitask R2",
    "table24_labelled_hybrid_physics_final_comparison.csv": "Table 24. Labelled, hybrid, and physics final comparison",
    "table25_best_overall_model_per_trait.csv": "Table 25. Best overall model per trait",
}


FIGURE_TITLES = {
    "fig1_selected_wavelengths.png": "Figure 1. Selected wavelengths for primary traits",
    "fig1_selected_wavelengths_annotated.png": "Figure 1b. Selected wavelengths annotated by vegetation absorption region",
    "fig2_r2_comparison.png": "Figure 2. R2 comparison",
    "fig3_rmse_comparison.png": "Figure 3. RMSE comparison",
    "fig4_predicted_vs_actual_cab.png": "Figure 4. Predicted vs actual Cab",
    "fig5_predicted_vs_actual_cw.png": "Figure 5. Predicted vs actual Cw",
    "fig6_predicted_vs_actual_cm.png": "Figure 6. Predicted vs actual Cm",
    "final_model_mean_r2_leaderboard.png": "Figure. Final model mean R2 leaderboard",
    "final_best_family_by_trait_r2.png": "Figure. Best labelled, hybrid, and physics-guided model families by trait",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=60, bottom=60, start=80, end=80) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for tag, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_landscape_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)


def add_portrait_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)


def clean_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        if abs(value) >= 1000 or (abs(value) > 0 and abs(value) < 0.001):
            return f"{value:.3e}"
        return f"{value:.4f}".rstrip("0").rstrip(".")
    text = str(value)
    if len(text) > 260:
        return text[:257] + "..."
    return text


def add_csv_table(doc: Document, csv_path: Path) -> None:
    df = pd.read_csv(csv_path)
    title = TABLE_TITLES.get(csv_path.name, csv_path.stem)
    doc.add_heading(title, level=2)
    meta = doc.add_paragraph(f"Source file: {csv_path.name}; rows: {len(df)}, columns: {len(df.columns)}.")
    meta.runs[0].italic = True
    meta.runs[0].font.size = Pt(8)

    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_cell_margins(table)

    header_cells = table.rows[0].cells
    for idx, col in enumerate(df.columns):
        cell = header_cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "E9EEF7")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(col))
        run.bold = True
        run.font.size = Pt(6.2)

    for _, row in df.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(df.columns):
            cell = cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            text = clean_value(row[col])
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(5.8 if len(df.columns) >= 12 else 6.8)

    doc.add_paragraph()


def add_image(doc: Document, image_path: Path, title: str, width: float = 9.4) -> None:
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    caption = doc.add_paragraph(f"Source file: {image_path.name}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)


def figure_title(path: Path) -> str:
    if path.name in FIGURE_TITLES:
        return FIGURE_TITLES[path.name]
    name = path.stem.replace("_", " ")
    return "Figure. " + name[:1].upper() + name[1:]


def main() -> None:
    build_docx(
        SOURCE,
        OUTPUT,
        "Detailed Project Report",
        "Band-efficient hyperspectral plant trait regression",
    )
    doc = Document(OUTPUT)

    add_landscape_section(doc)
    doc.add_heading("Appendix A. Complete Paper Tables", level=1)
    intro = doc.add_paragraph(
        "This appendix embeds the generated CSV paper tables directly in the Word report. "
        "Wide tables use small text so the complete columns remain present in the document."
    )
    intro.runs[0].font.size = Pt(9)
    for csv_path in sorted(TABLE_DIR.glob("*.csv")):
        add_csv_table(doc, csv_path)

    add_landscape_section(doc)
    doc.add_heading("Appendix B. Complete Figures", level=1)
    doc.add_paragraph(
        "This appendix embeds all generated paper figures and supporting result figures."
    )
    seen: set[Path] = set()
    figure_paths = list(sorted(PAPER_FIG_DIR.glob("*.png"))) + list(sorted(FIG_DIR.glob("*.png")))
    for image_path in figure_paths:
        resolved = image_path.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        add_image(doc, image_path, figure_title(image_path))

    doc.save(OUTPUT)
    doc.save(COMPLETE_OUTPUT)
    print(OUTPUT)
    print(COMPLETE_OUTPUT)


if __name__ == "__main__":
    main()
